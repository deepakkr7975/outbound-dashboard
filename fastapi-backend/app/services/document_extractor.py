"""
Document text extraction for AI-sequence generation.

Turns an uploaded reference document (PDF / DOCX / TXT / MD) into plain text so
the AI-sequence pipeline can use it as context. Kept deliberately small and
dependency-light: pypdf for PDFs, python-docx for .docx, and a decode for plain
text. Legacy binary .doc is not supported (needs external tooling); callers get
a clear UnsupportedFileType instead of garbled output.
"""

import io
import logging
import os
from typing import Iterable

logger = logging.getLogger(__name__)

# Extensions we can extract text from. Mirrors the frontend's `accept` list,
# minus legacy .doc which can't be parsed reliably without external tools.
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}

# Guard against oversized uploads and against documents that extract to almost
# nothing (scanned/image-only PDFs, empty files).
MAX_FILE_BYTES = 10 * 1024 * 1024  # 10 MB
MIN_TEXT_CHARS = 20
# Cap how much text we forward to the model to keep prompts bounded.
MAX_TEXT_CHARS = 20_000


class UnsupportedFileType(Exception):
    """The uploaded file's extension isn't one we can extract text from."""


class ExtractionError(Exception):
    """The file type is supported but text could not be extracted from it."""


def _ext(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower()


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
    except (PdfReadError, Exception) as e:  # noqa: BLE001 — corrupt/encrypted PDF
        raise ExtractionError(f"Could not read PDF: {e}") from e

    if reader.is_encrypted:
        # Try an empty-password decrypt; many PDFs are "encrypted" with no password.
        try:
            reader.decrypt("")
        except Exception as e:  # noqa: BLE001
            raise ExtractionError(
                "PDF is password-protected; remove the password and retry."
            ) from e

    parts: list[str] = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception as e:  # noqa: BLE001 — skip unreadable pages, keep going
            logger.warning("Skipping unreadable PDF page: %s", e)
    return "\n".join(parts)


def _extract_docx(data: bytes) -> str:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        doc = Document(io.BytesIO(data))
    except PackageNotFoundError as e:
        # Most commonly a legacy .doc mislabelled as .docx.
        raise ExtractionError(
            "File is not a valid .docx (legacy .doc is not supported — "
            "save it as .docx or PDF)."
        ) from e
    except Exception as e:  # noqa: BLE001
        raise ExtractionError(f"Could not read DOCX: {e}") from e

    lines: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def _extract_text_plain(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Could not decode text file (unknown encoding).")


def extract_text(filename: str, data: bytes) -> str:
    """
    Extract plain text from an uploaded document.

    Raises UnsupportedFileType for extensions we don't handle, and
    ExtractionError when a supported file can't be parsed or yields no usable
    text. The result is stripped and capped at MAX_TEXT_CHARS.
    """
    if not data:
        raise ExtractionError("Uploaded file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise ExtractionError(
            f"File is too large ({len(data) // (1024 * 1024)} MB); "
            f"limit is {MAX_FILE_BYTES // (1024 * 1024)} MB."
        )

    ext = _ext(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(
            f"Unsupported file type '{ext or filename}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}."
        )

    if ext == ".pdf":
        text = _extract_pdf(data)
    elif ext == ".docx":
        text = _extract_docx(data)
    else:  # .txt, .md
        text = _extract_text_plain(data)

    text = _collapse_blank_lines(text).strip()
    if len(text) < MIN_TEXT_CHARS:
        raise ExtractionError(
            "Could not extract meaningful text from the document "
            "(it may be empty or image-only/scanned)."
        )
    return text[:MAX_TEXT_CHARS]


def _collapse_blank_lines(text: str) -> str:
    lines: Iterable[str] = (line.rstrip() for line in text.splitlines())
    out: list[str] = []
    blank_run = 0
    for line in lines:
        if line.strip():
            blank_run = 0
            out.append(line)
        else:
            blank_run += 1
            if blank_run <= 1:  # keep single blank lines, drop runs
                out.append("")
    return "\n".join(out)
